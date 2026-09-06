import io
import json
import os
import re
import smtplib
import ssl
import unicodedata
import zipfile
from copy import deepcopy
from email.message import EmailMessage
from http.server import BaseHTTPRequestHandler

from docx import Document
from docx.text.paragraph import Paragraph

from templates_data import get_template

PT_MONTHS = ['JANEIRO','FEVEREIRO','MARÇO','ABRIL','MAIO','JUNHO','JULHO','AGOSTO','SETEMBRO','OUTUBRO','NOVEMBRO','DEZEMBRO']

def _template_bytes(kind):
    key={'atestado':'ATESTADO','acompanhante':'ACOMPANHANTE','simples':'SIMPLES','especial':'ESPECIAL'}[kind]
    return get_template(key)

def _br_date(iso):
    y,m,d = [int(x) for x in str(iso).split('-')]
    return f'{d:02d}/{m:02d}/{y:04d}'

def _ext_date(iso):
    y,m,d = [int(x) for x in str(iso).split('-')]
    return f'{d:02d} DE {PT_MONTHS[m-1]} DE {y}'

def _num_words(n):
    n=int(n)
    ones=['ZERO','UM','DOIS','TRÊS','QUATRO','CINCO','SEIS','SETE','OITO','NOVE']
    teens={10:'DEZ',11:'ONZE',12:'DOZE',13:'TREZE',14:'QUATORZE',15:'QUINZE',16:'DEZESSEIS',17:'DEZESSETE',18:'DEZOITO',19:'DEZENOVE'}
    tens={20:'VINTE',30:'TRINTA',40:'QUARENTA',50:'CINQUENTA',60:'SESSENTA',70:'SETENTA',80:'OITENTA',90:'NOVENTA'}
    if n < 10: return ones[n]
    if n < 20: return teens[n]
    if n < 100:
        t=(n//10)*10; r=n%10
        return tens[t] + (f' E {ones[r]}' if r else '')
    return str(n)

def _safe_name(s):
    s=unicodedata.normalize('NFD',str(s)).encode('ascii','ignore').decode()
    return re.sub(r'[^A-Za-z0-9]+','_',s).strip('_').upper() or 'PACIENTE'

def _parse_med(posologia):
    m=re.match(r'^(.*?)-{3,}\s*(.+)$',str(posologia or '').strip())
    return (m.group(1).strip(),m.group(2).strip()) if m else (str(posologia or '').strip(),'')

def _remove_p(p):
    el=p._element; el.getparent().remove(el); p._p=p._element=None

def _set_single_run_text(p,text):
    if not p.runs: p.add_run(text)
    else:
        p.runs[0].text=text
        for r in p.runs[1:]: r.text=''

def _load_doc(kind):
    return Document(io.BytesIO(_template_bytes(kind)))

def _save_doc(doc):
    out=io.BytesIO(); doc.save(out); return out.getvalue()

def generate_atestado(d):
    doc=_load_doc('atestado'); p=doc.paragraphs; r=p[7].runs
    r[1].text=d['paciente'].upper()
    r[3].text='PROCEDIMENTO DE '+d['procedimento'].upper()
    r[5].text=_br_date(d['dataProcedimento'])
    r[7].text=f"{d['diasAfastamento']} ({_num_words(d['diasAfastamento'])}) DIAS"
    r[9].text=_br_date(d['dataInternacao'])
    p[9].runs[0].text='CID-10: '+d['cid10'].upper()
    p[15].runs[0].text='FORTALEZA, '+_ext_date(d['hoje'])
    return _save_doc(doc)

def generate_acompanhante(d):
    doc=_load_doc('acompanhante'); p=doc.paragraphs; r=p[7].runs
    r[1].text=d['acompanhanteNome'].upper()
    r[2].text=(f" CPF: {d['acompanhanteCpf']} " if d.get('acompanhanteCpf') else ' ')
    r[3].text='esteve como acompanhante do paciente '
    r[4].text=d['paciente'].upper()
    r[6].text=f"{_br_date(d['inicio'])} a {_br_date(d['fim'])}"
    p[14].runs[0].text='Fortaleza-CE, '+_ext_date(d['hoje'])+'.'
    return _save_doc(doc)

def _fill_rx_cell(cell,d,via,special=False):
    ps=cell.paragraphs
    if special:
        _set_single_run_text(ps[3],f'Receituário de Controle Especial - {via}')
        _set_single_run_text(ps[5],'PACIENTE: '+d['paciente'].upper())
        _set_single_run_text(ps[6],'LOGRADOURO: '+str(d.get('endereco','')).upper())
        top_tpl=deepcopy(ps[10]._p); inst_tpl=deepcopy(ps[11]._p); blank_tpl=deepcopy(ps[12]._p)
        anchor=ps[14]
        for idx in [13,12,11,10]: _remove_p(ps[idx])
    else:
        _set_single_run_text(ps[5],'NOME: '+d['paciente'].upper())
        _set_single_run_text(ps[6],f'RECEITUÁRIO - {via}')
        top_tpl=deepcopy(ps[10]._p); inst_tpl=deepcopy(ps[11]._p); blank_tpl=deepcopy(ps[12]._p)
        anchor=ps[22]
        for idx in range(21,9,-1): _remove_p(ps[idx])
    for x in d.get('meds',[]):
        inst,qty=_parse_med(x.get('posologia',''))
        for tpl,text in [
            (top_tpl,f"{x.get('nome','').upper()} {'-'*35} {qty}".rstrip()),
            (inst_tpl,inst.upper()),
            (blank_tpl,'')
        ]:
            new_el=deepcopy(tpl); anchor._p.addprevious(new_el); np=Paragraph(new_el,anchor._parent)
            if text: _set_single_run_text(np,text)
            else:
                for rr in np.runs: rr.text=''
    outros=str(d.get('outros','')).strip()
    if outros:
        new_el=deepcopy(inst_tpl); anchor._p.addprevious(new_el); np=Paragraph(new_el,anchor._parent); _set_single_run_text(np,outros.upper())
        anchor._p.addprevious(deepcopy(blank_tpl))
    if special:
        _set_single_run_text(anchor,'Fortaleza-CE, '+_br_date(d['hoje']))
    else:
        _set_single_run_text(anchor,'FORTALEZA - CE')
        paras=cell.paragraphs
        for i,p in enumerate(paras):
            if p._p is anchor._p and i+1<len(paras):
                _set_single_run_text(paras[i+1],_br_date(d['hoje'])); break

def generate_receita(d,special=False):
    doc=_load_doc('especial' if special else 'simples'); table=doc.tables[0]
    _fill_rx_cell(table.cell(0,0),d,'1ª via',special)
    _fill_rx_cell(table.cell(0,1),d,'2ª via',special)
    return _save_doc(doc)

def validate_patient(p):
    if not isinstance(p,dict) or not str(p.get('paciente','')).strip(): return 'Paciente sem nome'
    if p.get('atestado'):
        a=p['atestado']
        if not all(str(a.get(k,'')).strip() for k in ['procedimento','cid10','dataProcedimento','dataInternacao','diasAfastamento']):
            return f"Atestado incompleto para {p['paciente']}"
    if p.get('acompanhante'):
        a=p['acompanhante']
        if not all(str(a.get(k,'')).strip() for k in ['acompanhanteNome','inicio','fim']):
            return f"Declaração de acompanhante incompleta para {p['paciente']}"
    if p.get('receitaEspecial') and not str(p['receitaEspecial'].get('endereco','')).strip():
        return f"Informe o logradouro para o receituário de controle especial de {p['paciente']}"
    return None

def normalize_request(body):
    if not isinstance(body,dict): return None,None
    if isinstance(body.get('payload'),dict):
        payload=body['payload']; recipient=body.get('recipient') or payload.get('recipient')
    else:
        payload={'hoje':body.get('hoje'),'pacientes':body.get('pacientes')}; recipient=body.get('recipient')
    return payload,recipient

def make_zip(payload):
    buf=io.BytesIO(); generated=0
    with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as z:
        for p in payload['pacientes']:
            base={'paciente':p['paciente'],'hoje':payload['hoje']}; s=_safe_name(p['paciente'])
            if p.get('atestado'):
                z.writestr(f'Atestado_{s}.docx',generate_atestado({**p['atestado'],**base})); generated+=1
            if p.get('acompanhante'):
                z.writestr(f'Declaracao_Acompanhante_{s}.docx',generate_acompanhante({**p['acompanhante'],**base})); generated+=1
            if p.get('receitaSimples'):
                z.writestr(f'Receituario_Simples_{s}.docx',generate_receita({**p['receitaSimples'],**base},False)); generated+=1
            if p.get('receitaEspecial'):
                z.writestr(f'Receituario_Controle_Especial_{s}.docx',generate_receita({**p['receitaEspecial'],**base},True)); generated+=1
    return buf.getvalue(),generated

def recipients(mode):
    a=os.getenv('PERSONAL_EMAIL'); b=os.getenv('SERVICE_EMAIL')
    if mode=='personal': return [a] if a else []
    if mode=='service': return [b] if b else []
    if mode=='both': return list(dict.fromkeys(x for x in [a,b] if x))
    return []

def send_zip(zbytes,to):
    host=os.getenv('SMTP_HOST'); port=int(os.getenv('SMTP_PORT','587')); user=os.getenv('SMTP_USER'); password=os.getenv('SMTP_PASS')
    if not all([host,user,password]): raise RuntimeError('Servidor de e-mail ainda não configurado')
    msg=EmailMessage(); msg['Subject']='Documentos de Alta CCP'; msg['From']=os.getenv('SMTP_FROM') or user; msg['To']=', '.join(to)
    msg.set_content('Documentos de alta gerados pelo aplicativo CCP. Os dados do paciente estão apenas nos arquivos anexos.')
    msg.add_attachment(zbytes,maintype='application',subtype='zip',filename='Documentos_Alta_CCP.zip')
    if port==465:
        with smtplib.SMTP_SSL(host,port,context=ssl.create_default_context(),timeout=25) as smtp:
            smtp.login(user,password); smtp.send_message(msg)
    else:
        with smtplib.SMTP(host,port,timeout=25) as smtp:
            smtp.ehlo(); smtp.starttls(context=ssl.create_default_context()); smtp.ehlo(); smtp.login(user,password); smtp.send_message(msg)

def response(handler,status,data):
    raw=json.dumps(data,ensure_ascii=False).encode('utf-8')
    handler.send_response(status); handler.send_header('Content-Type','application/json; charset=utf-8'); handler.send_header('Cache-Control','no-store'); handler.send_header('Content-Length',str(len(raw))); handler.end_headers(); handler.wfile.write(raw)

class handler(BaseHTTPRequestHandler):
    def do_GET(self): response(self,405,{'ok':False,'error':'Método não permitido'})
    def do_POST(self):
        try:
            if not os.getenv('APP_KEY') or self.headers.get('X-App-Key') != os.getenv('APP_KEY'):
                return response(self,401,{'ok':False,'error':'Chave do aplicativo inválida'})
            length=int(self.headers.get('Content-Length','0') or 0)
            body=json.loads(self.rfile.read(length).decode('utf-8')) if length else {}
            payload,recipient=normalize_request(body)
            if not payload or not isinstance(payload.get('pacientes'),list) or not payload['pacientes'] or len(payload['pacientes'])>20 or not payload.get('hoje'):
                return response(self,400,{'ok':False,'error':'Lote inválido'})
            for p in payload['pacientes']:
                err=validate_patient(p)
                if err:return response(self,400,{'ok':False,'error':err})
            to=recipients(recipient)
            if not to:return response(self,500,{'ok':False,'error':'E-mail de destino ainda não configurado no servidor'})
            zbytes,generated=make_zip(payload); send_zip(zbytes,to)
            return response(self,200,{'ok':True,'generated':generated,'patients':len(payload['pacientes']),'recipients':len(to)})
        except json.JSONDecodeError:
            return response(self,400,{'ok':False,'error':'JSON inválido'})
        except Exception:
            return response(self,500,{'ok':False,'error':'Falha ao gerar ou enviar os documentos'})
